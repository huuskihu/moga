# -*- coding: utf-8 -*-
"""
生成 presentation 报告：三张分图 + 一份 PPTX（通俗 × 专业 双表述）。
运行：python make_report.py
产物：
  results/fig_fit.png, fig_loss.png, fig_ga.png
  T-World_GD_vs_GA_报告.pptx
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

import config
from data import load_ground_truth
from model import fhn_simulate

BASE = config.BASE_DIR
RES = config.RESULTS_DIR


def load(key):
    return np.load(os.path.join(RES, f"{key}.npz"))


# ============ 数据 ============
t, y, y_clean, theta_true, n_spikes = load_ground_truth()
gd = load("gd_bad")
ga = load("ga")
neural = load("neural")
V_ga_best = fhn_simulate(ga["best_theta"], t, config)

# ============ 图1：拟合对比 ============
fig, ax = plt.subplots(figsize=(9, 4.2))
ax.plot(t, y_clean, color="0.35", lw=2, label="ground truth (4 次发放)")
ax.plot(t, gd["final_V"], color="#d62728", lw=1.2, alpha=0.9,
        label="GD + 白盒 ODE（28 次发放，卡死）")
ax.plot(t, V_ga_best, color="#2ca02c", ls="--", lw=1.6,
        label="遗传算法 + 白盒 ODE（恢复真值）")
ax.plot(t, neural["final_V"], color="#9467bd", ls="-.", lw=1.3,
        label="GD + 神经代理（平滑拟合）")
ax.set_xlabel("t")
ax.set_ylabel("V(t)")
ax.set_title("拟合对比：同一份数据，三种训练方式")
ax.legend(fontsize=8, loc="upper right")
fig.tight_layout()
fig.savefig(os.path.join(RES, "fig_fit.png"), dpi=130)
plt.close(fig)

# ============ 图2：损失曲线 ============
fig, ax = plt.subplots(figsize=(9, 4.2))
g_hist = gd["loss_history"]
g_fin = np.isfinite(g_hist)
ax.plot(np.arange(len(g_hist))[g_fin], g_hist[g_fin], color="#d62728", lw=1.4,
        label="GD + 白盒 ODE（卡死在 2.06，不下降）")
n_hist = neural["loss_history"]
ax.plot(np.arange(len(n_hist)), n_hist, color="#9467bd", lw=1.4,
        label="GD + 神经代理（平滑降到 0.06）")
ax.axhline(0.00043, color="#2ca02c", ls=":", lw=1.2,
           label="遗传算法最优 MSE = 0.00043")
ax.set_yscale("log")
ax.set_xlabel("epoch")
ax.set_ylabel("loss (MSE, log)")
ax.set_title("损失曲线：GD 对刚性白盒 ODE 卡死，对可微代理平滑收敛")
ax.legend(fontsize=8)
fig.tight_layout()
fig.savefig(os.path.join(RES, "fig_loss.png"), dpi=130)
plt.close(fig)

# ============ 图3：遗传算法收敛 ============
fig, ax = plt.subplots(figsize=(9, 4.2))
ax.plot(ga["history_f1"], color="#2ca02c", lw=1.5, marker="o", ms=3,
        label="最小 MSE (f1)")
ax.set_xlabel("代数 (generation)")
ax.set_ylabel("最小 MSE (f1)")
ax.set_yscale("log")
ax.set_title("遗传算法收敛：MSE 快速降至 ~0.0004，100% 个体发放正确")
ax.legend(fontsize=9)
ax.set_ylim(bottom=1e-4)
fig.tight_layout()
fig.savefig(os.path.join(RES, "fig_ga.png"), dpi=130)
plt.close(fig)

print("三张分图已生成。")

# ============ PPT ============
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x20, 0x20, 0x20)
GRAY = RGBColor(0x59, 0x59, 0x59)
PLAIN_C = RGBColor(0x2E, 0x74, 0xB5)   # 通俗（蓝）
PRO_C = RGBColor(0x70, 0x70, 0x70)      # 专业（灰）

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def set_title(slide, text, size=26, color=ACCENT):
    slide.shapes.title.text = text
    tf = slide.shapes.title.text_frame
    tf.paragraphs[0].font.size = Pt(size)
    tf.paragraphs[0].font.color.rgb = color
    tf.paragraphs[0].font.bold = True


def add_dual(body, plain, prof, first=False):
    """追加一组「通俗 + 专业」双表述。plain 为通俗，prof 为专业。"""
    if first:
        p0 = body.paragraphs[0]
        p1 = body.add_paragraph()
    else:
        p0 = body.add_paragraph()
        p1 = body.add_paragraph()
    p0.text = "通俗 · " + plain
    p0.level = 0
    p0.font.size = Pt(16)
    p0.font.color.rgb = PLAIN_C
    p0.space_before = Pt(6)
    p1.text = "专业 · " + prof
    p1.level = 1
    p1.font.size = Pt(13)
    p1.font.color.rgb = PRO_C
    p1.space_after = Pt(6)


def content_slide(title, pairs):
    """pairs: list of (plain, prof)"""
    s = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(s, title)
    body = s.placeholders[1].text_frame
    body.word_wrap = True
    for i, (plain, prof) in enumerate(pairs):
        add_dual(body, plain, prof, first=(i == 0))
    return s


def image_slide(title, img_path, plain, prof):
    s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    set_title(s, title)
    left = (prs.slide_width - Inches(11.0)) // 2
    s.shapes.add_picture(img_path, left, Inches(1.5), width=Inches(11.0))
    tb = s.shapes.add_textbox(Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "通俗 · " + plain
    p0.font.size = Pt(15)
    p0.font.color.rgb = PLAIN_C
    p1 = tf.add_paragraph()
    p1.text = "专业 · " + prof
    p1.font.size = Pt(12)
    p1.font.color.rgb = PRO_C
    return s


# --- 封面 ---
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "梯度下降  vs  遗传算法"
s.shapes.title.text_frame.paragraphs[0].font.size = Pt(42)
s.shapes.title.text_frame.paragraphs[0].font.bold = True
s.shapes.title.text_frame.paragraphs[0].font.color.rgb = ACCENT
sub = s.placeholders[1]
sub.text = ("通俗 · 给模型调参，到底该“算梯度下山”，还是该“进化式海选”？\n"
            "专业 · 基于梯度的优化与多目标进化优化，在白盒 ODE 参数标定任务中的对照研究")
for p in sub.text_frame.paragraphs:
    p.font.size = Pt(16)
    p.font.color.rgb = GRAY

# --- 问题 ---
content_slide("问题：两种 AI4Science 模型，两种训练范式", [
    ("两种 AI 模型都在“让计算机自动拟合数据”，只是调参方式完全不同。",
     "二者均属数据驱动的参数估计任务，却分别采用梯度反向传播与多目标进化优化两种优化范式。"),
    ("T-World 是把心肌细胞写成“每个参数都看得懂含义”的物理方程组，所以用遗传算法来调参。",
     "T-World 为机理驱动的白盒常微分方程（ODE）模型，参数具生理学可解释性，故采用多目标遗传算法（MOGA）标定。"),
    ("ProteoFlow / UMA 把规律装进“光滑可导的神经网络”，所以能直接求梯度。",
     "ProteoFlow / UMA 采用可微神经网络参数化（流匹配速度场 / 等变图网络），支持端到端基于梯度的训练。"),
    ("核心疑问：同样拟合数据，凭啥一个用进化、一个用梯度？",
     "核心问题：同一拟合任务，其优化范式的选择差异，其内在决定因素为何？"),
])

# --- 方法 ---
content_slide("方法：一个“麻雀虽小、五脏俱全”的替身模型", [
    ("不搬原模型（太大、要超算），造了个保留关键特征的迷你细胞模型。",
     "采用 FitzHugh–Nagumo 可兴奋系统作为降维替身，保留“白盒 / 刚性 / 阈值”三大核心特征。"),
    ("它一边动得快、一边动得慢（叫刚性）；发不发放是“一刀切”的（有阈值）。",
     "eps << 1 引入刚性，快慢时间尺度相差约 50 倍；系统具兴奋阈值，发放与否为二值事件，类比早/延迟后除极（EAD/DAD）。"),
    ("有 4 个旋钮参数，真实值能让细胞“每刺激一下发一次”。",
     "4 个待拟合参数 [eps, a, b, amp]，真实值 [0.02, 0.8, 0.8, 0.8] 对应 1:1 刺激-发放；目标为向量 [f_MSE, f_fire]。"),
])

# --- 对比设计 ---
content_slide("对比设计：只换“模型长相”，其余全不动", [
    ("做一次控制变量：数据、优化器都不动，只换模型的参数化方式。",
     "以参数化方式为唯一控制变量，保持数据与优化器不变，进行受控对照实验。"),
    ("路线①：硬要用梯度下降的白盒 ODE（T-World 的反例）。",
     "方案①：GD + 白盒 ODE——对应“若 T-World 坚持梯度下降”的反事实情形。"),
    ("路线②：进化算法 + 白盒 ODE（T-World 实际做法）。",
     "方案②：MOGA + 白盒 ODE——对应 T-World 实际采用的标定方式。"),
    ("路线③：梯度下降 + 可微神经代理（ProteoFlow/UMA 范式）。",
     "方案③：GD + 可微神经代理——对应 ProteoFlow / UMA 的可微参数化范式。"),
])

# --- 结果图1 ---
image_slide("结果一：拟合效果",
            os.path.join(RES, "fig_fit.png"),
            "一眼看出：红色（梯度下降）卡死、发错了 28 次；绿色（遗传算法）几乎和真值重合。",
            "梯度下降陷入错误发放模式的局部极小（28 次 vs 4 次）；遗传算法恢复真实波形（MSE 0.0004）。")

# --- 结果图2 ---
image_slide("结果二：损失曲线",
            os.path.join(RES, "fig_loss.png"),
            "损失掉到 2.06 就再也不动了——不是没跑够，是真下不去了；紫色（神经代理）却一路光滑向下。",
            "GD+白盒 ODE 收敛至非零平台（局部极小）；GD+神经代理单调下降至 0.06；GA 达 0.0004。")

# --- 结果图3 ---
image_slide("结果三：遗传算法收敛",
            os.path.join(RES, "fig_ga.png"),
            "几十代就找到答案，100% 的候选都“发对了次数”。",
            "MOGA 在数十代内收敛至 MSE≈0.0004，末代种群 100% 满足发放次数约束（f_fire=0）。")

# --- 结论 ---
content_slide("结论：差异在“参数化”，不在“优化器”", [
    ("工具没坏——梯度下降对可微的神经代理，0.5 秒就收敛好了。",
     "梯度下降本身有效：对可微神经代理，损失平滑收敛（MSE 0.06，0.5s），证明失效并非源于优化器。"),
    ("是模型“长得太硬”，让梯度下降使不上劲：卡在错误放电，还守不住参数下界。",
     "刚性白盒 ODE 的病态条件数与非光滑解使 GD 陷入局部极小，并逸出可行域（eps 略低于生理下界 0.01）。"),
    ("进化算法不靠导数，稳，找回真值。",
     "MOGA 不依赖梯度，对刚性、阈值与不可导目标鲁棒，稳健恢复真值（MSE 0.0004）。"),
    ("给你选型一句话：要保真可解释 → 用进化；要端到端又快 → 用可微 + 梯度。",
     "工程准则：机理保真优先时选进化优化；端到端可训练性优先时选可微参数化 + 梯度下降。"),
])

out = os.path.join(BASE, "T-World_GD_vs_GA_报告.pptx")
prs.save(out)
print("PPT 已生成：", out)